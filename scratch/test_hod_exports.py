import io
import pytest
import openpyxl
from app import app


@pytest.fixture
def client():
    app.config["TESTING"] = True
    app.config["SECRET_KEY"] = "test-secret"
    with app.test_client() as client:
        yield client


def test_unauthenticated_export_blocked(client):
    """Unauthenticated users calling /api/ endpoints must be blocked with 401 Unauthorized."""
    endpoints = [
        "/api/hod/export/students/csv",
        "/api/hod/export/students/xlsx",
        "/api/hod/export/faculty/csv",
        "/api/hod/export/faculty/xlsx",
        "/api/hod/export/summary/csv",
        "/api/hod/export/summary/xlsx"
    ]
    for ep in endpoints:
        res = client.get(ep)
        assert res.status_code == 401
        data = res.get_json()
        assert data["error"] == "Unauthorized"


def test_hod_exports_success(client):
    """Authenticated HOD (CSE) can export all CSV and Excel reports."""
    # Login as CSE HOD
    login_res = client.post("/login", data={"username": "hod_cse", "password": "Hod@123"}, follow_redirects=False)
    assert login_res.status_code == 302
    assert "/hod" in login_res.headers["Location"]

    # 1. Students CSV
    res = client.get("/api/hod/export/students/csv")
    assert res.status_code == 200
    assert "text/csv" in res.headers["Content-Type"]
    assert "attachment" in res.headers["Content-Disposition"]
    assert "CSE_Student_Report" in res.headers["Content-Disposition"]
    csv_text = res.data.decode("utf-8")
    assert "Student ID,Full Name,Department" in csv_text
    assert "CSE" in csv_text
    # Ensure all student rows in export belong to CSE
    lines = csv_text.strip().split("\n")[1:]  # skip header
    for line in lines:
        if line.strip():
            cols = line.split(",")
            assert cols[2].strip() == "CSE"

    # 2. Students Excel
    res = client.get("/api/hod/export/students/xlsx")
    assert res.status_code == 200
    assert "openxmlformats-officedocument" in res.headers["Content-Type"]
    assert "attachment" in res.headers["Content-Disposition"]
    assert "CSE_Student_Report" in res.headers["Content-Disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    assert "CSE Students" in wb.sheetnames
    ws = wb["CSE Students"]
    assert "Student Cohort Performance" in ws.cell(row=1, column=1).value

    # 3. Faculty CSV
    res = client.get("/api/hod/export/faculty/csv")
    assert res.status_code == 200
    assert "text/csv" in res.headers["Content-Type"]
    assert "CSE_Faculty_Report" in res.headers["Content-Disposition"]
    fac_csv = res.data.decode("utf-8")
    assert "Faculty ID,Full Name,Designation,Department" in fac_csv
    assert "Attendance tracking not active in current system" in fac_csv

    # 4. Faculty Excel
    res = client.get("/api/hod/export/faculty/xlsx")
    assert res.status_code == 200
    assert "CSE_Faculty_Report" in res.headers["Content-Disposition"]
    wb_fac = openpyxl.load_workbook(io.BytesIO(res.data))
    assert "CSE Faculty" in wb_fac.sheetnames

    # 5. Summary CSV
    res = client.get("/api/hod/export/summary/csv")
    assert res.status_code == 200
    assert "CSE_Department_Summary" in res.headers["Content-Disposition"]
    summary_csv = res.data.decode("utf-8")
    assert "Metric Category,Metric Name,Value" in summary_csv
    assert "CSE" in summary_csv

    # 6. Summary Excel
    res = client.get("/api/hod/export/summary/xlsx")
    assert res.status_code == 200
    assert "CSE_Department_Summary" in res.headers["Content-Disposition"]
    wb_sum = openpyxl.load_workbook(io.BytesIO(res.data))
    assert "CSE Executive Summary" in wb_sum.sheetnames


def test_cross_role_and_dept_isolation(client):
    """Students and other non-HOD roles are blocked with 403 Forbidden; AIML HOD gets AIML data only."""
    # Student login (student001 / Student@123)
    login_res = client.post("/login", data={"username": "student001", "password": "Student@123"})
    assert login_res.status_code == 302
    res = client.get("/api/hod/export/students/csv")
    assert res.status_code == 403
    data = res.get_json()
    assert data["error"] == "Forbidden"
    client.get("/logout")

    # Faculty login (faculty001 / Faculty@123)
    login_res = client.post("/login", data={"username": "faculty001", "password": "Faculty@123"})
    assert login_res.status_code == 302
    res = client.get("/api/hod/export/students/csv")
    assert res.status_code == 403
    client.get("/logout")

    # Login as AIML HOD
    client.post("/login", data={"username": "hod_aiml", "password": "Hod@123"})
    res = client.get("/api/hod/export/students/csv")
    assert res.status_code == 200
    assert "AIML_Student_Report" in res.headers["Content-Disposition"]
    csv_text = res.data.decode("utf-8")
    lines = csv_text.strip().split("\n")[1:]
    for line in lines:
        if line.strip():
            cols = line.split(",")
            assert cols[2].strip() == "AIML"
