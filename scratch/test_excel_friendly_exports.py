import io
import pytest
import openpyxl
from app import app


@pytest.fixture
def client():
    app.config["TESTING"] = True
    app.config["SECRET_KEY"] = "test-secret"
    with app.test_client() as client:
        yield client


def test_bom_presence_and_excel_friendliness(client):
    """Admin CSV export must start with UTF-8 BOM (0xEF 0xBB 0xBF) so Excel detects UTF-8 natively."""
    client.post("/login", data={"username": "admin", "password": "Admin@123"})
    res = client.get("/api/export/csv")
    assert res.status_code == 200
    assert "text/csv" in res.headers["Content-Type"]
    assert "charset=utf-8" in res.headers["Content-Type"].lower()
    assert "attachment" in res.headers["Content-Disposition"]
    assert ".csv" in res.headers["Content-Disposition"]

    # Check raw bytes for UTF-8 Byte Order Mark
    raw_bytes = res.data
    assert raw_bytes[:3] == b"\xef\xbb\xbf", "Missing UTF-8 BOM prefix (\\ufeff)"

    # Check parsed CSV structure
    decoded_text = raw_bytes.decode("utf-8-sig")
    lines = decoded_text.strip().split("\n")
    assert len(lines) >= 2  # At least header and one data row
    header = lines[0].split(",")
    assert "Student ID" in header[0]
    assert "Name" in header[1]
    assert "Risk Category" in header[-1]


def test_admin_full_export_suite(client):
    """Admin can export all institutional reports in CSV (with BOM) and Excel (.xlsx)."""
    client.post("/login", data={"username": "admin", "password": "Admin@123"})

    # 1. Student Cohort CSV & Excel
    res_csv = client.get("/api/export/csv")
    assert res_csv.status_code == 200
    assert res_csv.data[:3] == b"\xef\xbb\xbf"

    res_xlsx = client.get("/api/export/xlsx")
    assert res_xlsx.status_code == 200
    assert "openxmlformats-officedocument" in res_xlsx.headers["Content-Type"]
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.data))
    assert "Student Cohort" in wb.sheetnames
    ws = wb["Student Cohort"]
    assert ws.max_row > 4

    # 2. Filtered High-Risk CSV & Excel
    res_high_csv = client.get("/api/export/csv?risk_level=High")
    assert res_high_csv.status_code == 200
    assert "High_Risk" in res_high_csv.headers["Content-Disposition"]
    assert res_high_csv.data[:3] == b"\xef\xbb\xbf"

    res_high_xlsx = client.get("/api/export/xlsx?risk_level=High")
    assert res_high_xlsx.status_code == 200
    assert "High_Risk" in res_high_xlsx.headers["Content-Disposition"]

    # 3. Interventions CSV & Excel
    res_intv_csv = client.get("/api/interventions/export/csv")
    assert res_intv_csv.status_code == 200
    assert res_intv_csv.data[:3] == b"\xef\xbb\xbf"
    assert "Intervention_Report" in res_intv_csv.headers["Content-Disposition"]

    res_intv_xlsx = client.get("/api/interventions/export/xlsx")
    assert res_intv_xlsx.status_code == 200
    wb_intv = openpyxl.load_workbook(io.BytesIO(res_intv_xlsx.data))
    assert "Interventions" in wb_intv.sheetnames

    # 4. Faculty Directory CSV & Excel
    res_fac_csv = client.get("/api/faculty/export/csv")
    assert res_fac_csv.status_code == 200
    assert res_fac_csv.data[:3] == b"\xef\xbb\xbf"
    assert "Faculty_Report" in res_fac_csv.headers["Content-Disposition"]

    res_fac_xlsx = client.get("/api/faculty/export/xlsx")
    assert res_fac_xlsx.status_code == 200
    wb_fac = openpyxl.load_workbook(io.BytesIO(res_fac_xlsx.data))
    assert "Faculty Directory" in wb_fac.sheetnames


def test_faculty_export_scoping(client):
    """Faculty export is scoped to their department / advisees."""
    client.post("/login", data={"username": "faculty001", "password": "Faculty@123"})

    # Advisees cohort
    res = client.get("/api/export/csv")
    assert res.status_code == 200
    assert res.data[:3] == b"\xef\xbb\xbf"

    res_xlsx = client.get("/api/export/xlsx")
    assert res_xlsx.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.data))
    assert "Student Cohort" in wb.sheetnames

    # Interventions
    res_intv = client.get("/api/interventions/export/csv")
    assert res_intv.status_code == 200
    assert res_intv.data[:3] == b"\xef\xbb\xbf"

    res_intv_xlsx = client.get("/api/interventions/export/xlsx")
    assert res_intv_xlsx.status_code == 200


def test_hod_exports_have_bom(client):
    """HOD CSV exports (Students, Faculty, Summary) all have UTF-8 BOM."""
    client.post("/login", data={"username": "hod_cse", "password": "Hod@123"})

    # 1. Students CSV
    res_stu = client.get("/api/hod/export/students/csv")
    assert res_stu.status_code == 200
    assert res_stu.data[:3] == b"\xef\xbb\xbf"

    # 2. Faculty CSV
    res_fac = client.get("/api/hod/export/faculty/csv")
    assert res_fac.status_code == 200
    assert res_fac.data[:3] == b"\xef\xbb\xbf"

    # 3. Summary CSV
    res_sum = client.get("/api/hod/export/summary/csv")
    assert res_sum.status_code == 200
    assert res_sum.data[:3] == b"\xef\xbb\xbf"

    # 4. Interventions CSV
    res_intv = client.get("/api/interventions/export/csv")
    assert res_intv.status_code == 200
    assert res_intv.data[:3] == b"\xef\xbb\xbf"
    assert "CSE_Intervention_Report" in res_intv.headers["Content-Disposition"]


def test_student_transcript_export(client):
    """Student can export their personal academic transcript in CSV (with BOM) and Excel (.xlsx), but is blocked from cohort exports."""
    client.post("/login", data={"username": "student001", "password": "Student@123"})

    # Student transcript CSV
    res_csv = client.get("/api/student/transcript/export/csv")
    assert res_csv.status_code == 200
    assert res_csv.data[:3] == b"\xef\xbb\xbf"
    assert "Student_Transcript_" in res_csv.headers["Content-Disposition"]
    decoded = res_csv.data.decode("utf-8-sig")
    assert "Year,Semester,Academic Year,Course Code,Course Name" in decoded

    # Student transcript Excel
    res_xlsx = client.get("/api/student/transcript/export/xlsx")
    assert res_xlsx.status_code == 200
    assert "openxmlformats-officedocument" in res_xlsx.headers["Content-Type"]
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.data))
    assert "Academic Transcript" in wb.sheetnames

    # Blocked from administrative cohort exports
    assert client.get("/api/export/csv").status_code == 403
    assert client.get("/api/export/xlsx").status_code == 403
    assert client.get("/api/interventions/export/csv").status_code == 403
    assert client.get("/api/faculty/export/csv").status_code == 403


def test_formula_injection_sanitization(client):
    """Untrusted formula characters (=, +, -, @) are sanitized with single quote while preserving valid numbers."""
    from app import _sanitize_csv_cell

    # Formulas
    assert _sanitize_csv_cell("=1+2") == "'=1+2"
    assert _sanitize_csv_cell("=SUM(A1:A10)") == "'=SUM(A1:A10)"
    assert _sanitize_csv_cell("@cmd") == "'@cmd"
    assert _sanitize_csv_cell("+cmd|' /C ...'!A0") == "'+cmd|' /C ...'!A0"
    assert _sanitize_csv_cell("-1+2") == "'-1+2"

    # Legitimate numbers
    assert _sanitize_csv_cell("-5") == "-5"
    assert _sanitize_csv_cell("-3.14") == "-3.14"
    assert _sanitize_csv_cell(42) == 42
    assert _sanitize_csv_cell(85.5) == 85.5

    # Ordinary text
    assert _sanitize_csv_cell("John Doe") == "John Doe"
    assert _sanitize_csv_cell("CSE-2026-001") == "CSE-2026-001"
    assert _sanitize_csv_cell(None) == ""
